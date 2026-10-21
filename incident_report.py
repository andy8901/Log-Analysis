"""
Incident report on the Neosky letterhead (Word .docx), built from report["investigation"] (see investigator.py).

Page 1-1.5: verdict box (HUMAN ERROR / SYSTEM ERROR), in simple words, summary, what went wrong, possible causes,
why exactly it went wrong (and why it is not the other kind of error), corrective actions.
Annexure: every flight figure with its healthy range and what it means, the 3D path and top-down map, coordinates,
timeline, health checks, charts with a sentence on how to read each, and the drone's safety settings.

The letterhead's own header and footer are kept untouched; only its body is replaced.
"""
import os
import tempfile
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.collections import LineCollection  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

LETTERHEAD = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "neosky_letterhead.docx")

NAVY, TEAL, MUTED, GREY = "0E3A4F", "1A6E8A", "55656E", "7A8A93"
SIDE = {"human": ("C4581B", "FBEBDD"), "system": ("5A4FCF", "E9E6FB"), "none": ("2E7D32", "E2F0E2"), "unknown": ("55656E", "EEF1F2")}
KIND_COLOR = {"normal": "2E7D32", "hard_landing": "9A5800", "crash": "B3261E", "log_end": "B3261E"}
SEV = {"crit": ("Problem", "B3261E"), "warn": ("Warning", "9A5800"), "ok": ("OK", "2E7D32"), "info": ("Info", "0F6F7D")}
ROLE_COLOR = {"cause": "B3261E", "trigger": "9A5800", "contributing": MUTED}

FS_ACT_RC = {0: "Disabled", 1: "RTL", 2: "Continue mission in Auto, else RTL", 3: "Land", 4: "SmartRTL, else RTL", 5: "SmartRTL, else Land", 6: "Auto landing sequence, else RTL", 7: "Brake, else Land"}
FS_ACT_GCS = {0: "Disabled", 1: "RTL", 2: "Continue mission in Auto, else RTL", 3: "SmartRTL, else RTL", 4: "SmartRTL, else Land", 5: "Land", 6: "Auto landing sequence, else RTL", 7: "Brake, else Land"}
FS_ACT_BAT = {0: "Warn only", 1: "Land", 2: "RTL", 3: "SmartRTL, else RTL", 4: "SmartRTL, else Land", 5: "Terminate (motors off)", 6: "Auto landing sequence, else RTL", 7: "Brake, else Land"}
FS_OPT = ["continue Auto on RC loss", "continue Auto on GCS loss", "continue Guided on RC loss", "continue landing on any failsafe", "continue pilot modes on GCS loss", "release gripper"]


def _num(v):
    return f"{v:g}" if isinstance(v, (int, float)) else str(v)


SETTINGS = [
    ("FS_THR_ENABLE", "If the RC signal is lost", lambda v: FS_ACT_RC.get(int(v), _num(v))),
    ("FS_GCS_ENABLE", "If telemetry (ground station) is lost", lambda v: FS_ACT_GCS.get(int(v), _num(v))),
    ("FS_GCS_TIMEOUT", "Telemetry loss timeout", lambda v: f"{v:g} s"),
    ("FS_OPTIONS", "Failsafe options", lambda v: ", ".join(o for i, o in enumerate(FS_OPT) if int(v) & (1 << i)) or "none"),
    ("BATT_FS_LOW_ACT", "When the battery is low", lambda v: FS_ACT_BAT.get(int(v), _num(v))),
    ("BATT_LOW_VOLT", "Low battery voltage", lambda v: f"{v:g} V"),
    ("BATT_FS_CRT_ACT", "When the battery is critical", lambda v: FS_ACT_BAT.get(int(v), _num(v))),
    ("BATT_CRT_VOLT", "Critical battery voltage", lambda v: f"{v:g} V"),
    ("BATT_CAPACITY", "Battery capacity", lambda v: f"{v:g} mAh"),
    ("FS_EKF_ACTION", "If navigation becomes unreliable", lambda v: {1: "Land", 2: "AltHold", 3: "Land even in Stabilize"}.get(int(v), _num(v))),
    ("FS_CRASH_CHECK", "Stop motors after a crash", lambda v: "On" if v else "Off"),
    ("RTL_ALT", "Return-home altitude", lambda v: f"{v / 100:g} m"),
    ("LAND_SPEED", "Final landing speed", lambda v: f"{v / 100:g} m/s"),
    ("FENCE_ENABLE", "Geofence", lambda v: "On" if v else "Off"),
    ("ARMING_CHECK", "Pre-arm safety checks", lambda v: "Off" if v == 0 else "All" if v == 1 else "Some"),
    ("ANGLE_MAX", "Maximum lean angle", lambda v: f"{v / 100:g}°"),
]


# --------------------------------------------------------------------------------------------- docx helpers
def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _borders(table, color="C9D1D6"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def _run(p, text, bold=False, italic=False, size=10, color=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _para(doc, after=3, before=0, align=None, indent=None, hang=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before = Pt(after), Pt(before)
    pf.line_spacing = 1.08
    if align:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Inches(indent)
        if hang:
            pf.first_line_indent = Inches(-hang)
    return p


def _section(doc, text):
    p = _para(doc, after=3, before=8)
    p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "2"), ("w:color", "54C8E8")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)
    _run(p, text.upper(), bold=True, size=10.5, color=NAVY)


def _table(doc, rows, widths, head=None, size=9):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for i, gc in enumerate(t._tbl.tblGrid.gridCol_lst):
        gc.w = Inches(widths[i])
    all_rows = ([head] if head else []) + rows
    for ri, row in enumerate(all_rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            c = cells[ci]
            c.width = Inches(widths[ci])
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            spec = val if isinstance(val, dict) else {"t": val}
            if head and ri == 0:
                _shade(c, "E4EAED")
                _run(p, str(spec["t"]), bold=True, size=size)
                continue
            if spec.get("fill"):
                _shade(c, spec["fill"])
            for part in spec.get("runs") or [spec]:
                _run(p, str(part.get("t", "")), bold=part.get("b", False), size=part.get("sz", size), color=part.get("color"), italic=part.get("i", False))
    _para(doc, after=2)
    return t


def _clear_body(doc):
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)


# --------------------------------------------------------------------------------------------- pictures
def _path_3d(inv, png):
    tr = [p for p in inv.get("track") or [] if p.get("armed")] or inv.get("track") or []
    if len(tr) < 2:
        return False
    fig = plt.figure(figsize=(7.5, 4.3), dpi=150)
    ax = fig.add_subplot(111, projection="3d")
    xs, ys, hs = [p["x"] for p in tr], [p["z"] for p in tr], [max(0, p["y"]) for p in tr]
    top = max(1, max(hs))
    cmap = plt.get_cmap("plasma")
    for i in range(1, len(tr)):
        ax.plot(xs[i - 1:i + 1], ys[i - 1:i + 1], hs[i - 1:i + 1], color=cmap(0.15 + 0.7 * hs[i] / top), linewidth=1.6)
    ax.plot(xs, ys, [0] * len(xs), color="#9aa7ae", linewidth=0.8, alpha=0.7)
    step = max(1, len(tr) // 45)
    for i in range(0, len(tr), step):
        if hs[i] > 2:
            ax.plot([xs[i], xs[i]], [ys[i], ys[i]], [0, hs[i]], color="#c7d0d5", linewidth=0.5)
    colors = {"accent": "#0f6f7d", "crit": "#b3261e", "warn": "#c77700", "ok": "#2e7d32"}
    for m in inv.get("markers") or []:
        ax.scatter([m["x"]], [m["z"]], [max(0, m["y"])], color=colors.get(m["color"], "#333"), s=28, depthshade=False)
        ax.text(m["x"], m["z"], max(0, m["y"]) + top * 0.06, m["label"], fontsize=7, color=colors.get(m["color"], "#333"))
    ax.set_xlabel("East (m)", fontsize=7)
    ax.set_ylabel("North (m)", fontsize=7)
    ax.set_zlabel("Height (m)", fontsize=7)
    ax.tick_params(labelsize=6)
    span = max(max(xs) - min(xs), max(ys) - min(ys), 10)
    cx, cy = (max(xs) + min(xs)) / 2, (max(ys) + min(ys)) / 2
    ax.set_xlim(cx - span / 2, cx + span / 2)
    ax.set_ylim(cy - span / 2, cy + span / 2)
    ax.set_zlim(0, top * 1.15)
    ax.view_init(elev=28, azim=-60)
    fig.tight_layout()
    fig.savefig(png)
    plt.close(fig)
    return True


def _path_top(inv, png):
    tr = [p for p in inv.get("track") or [] if p.get("armed")] or inv.get("track") or []
    if len(tr) < 2:
        return False
    fig, ax = plt.subplots(figsize=(4.6, 4.6), dpi=150)
    pts = [[p["x"], p["z"]] for p in tr]
    segs = [[pts[i - 1], pts[i]] for i in range(1, len(pts))]
    hs = [max(0, p["y"]) for p in tr[1:]]
    lc = LineCollection(segs, cmap="plasma", linewidth=1.5)
    lc.set_array(hs)
    ax.add_collection(lc)
    colors = {"accent": "#0f6f7d", "crit": "#b3261e", "warn": "#c77700", "ok": "#2e7d32"}
    for m in inv.get("markers") or []:
        ax.scatter([m["x"]], [m["z"]], color=colors.get(m["color"], "#333"), s=24, zorder=3)
        ax.annotate(m["label"], (m["x"], m["z"]), fontsize=6.5, xytext=(4, 4), textcoords="offset points")
    ax.set_aspect("equal", adjustable="datalim")
    ax.autoscale()
    ax.set_xlabel("East of take-off (m)", fontsize=7)
    ax.set_ylabel("North of take-off (m)", fontsize=7)
    ax.tick_params(labelsize=6)
    ax.grid(alpha=0.3)
    cb = fig.colorbar(lc, ax=ax, fraction=0.04, pad=0.02)
    cb.set_label("Height (m)", fontsize=7)
    cb.ax.tick_params(labelsize=6)
    fig.tight_layout()
    fig.savefig(png)
    plt.close(fig)
    return True


def _chart(png, t, series, ylabel, inv, floor=None):
    series = [s for s in series if s[1] and any(v is not None for v in s[1])]
    if not series:
        return False
    fig, ax = plt.subplots(figsize=(7.2, 2.3), dpi=150)
    for label, ys, style in series:
        ax.plot([x / 60 for x in t], [v if v is not None else float("nan") for v in ys], label=label, **{"linewidth": 1.1, **style})
    ch = inv.get("chart") or {}
    if ch.get("contact") is not None:
        ax.axvspan(ch["contact"] / 60, t[-1] / 60, color="#b3261e", alpha=0.08, lw=0)
    for fs in ch.get("failsafes") or []:
        ax.axvline(fs / 60, color="#c77700", linestyle=":", linewidth=1)
    if floor:
        ax.axhline(floor, color="#b3261e", linestyle="--", linewidth=0.9)
    ax.set_xlabel("Minutes since arming", fontsize=7)
    ax.set_ylabel(ylabel, fontsize=7)
    ax.tick_params(labelsize=6.5)
    ax.grid(alpha=0.3)
    if len(series) > 1:
        ax.legend(fontsize=6.5, loc="upper right", ncol=min(4, len(series)))
    fig.tight_layout()
    fig.savefig(png)
    plt.close(fig)
    return True


# --------------------------------------------------------------------------------------------- main
def generate_incident_report(report, mission, output_path, template_path=LETTERHEAD):
    inv = report.get("investigation")
    if not inv:
        raise ValueError("This log has no flight investigation (only .bin dataflash logs are supported).")
    R = inv["report"]
    doc = Document(template_path) if os.path.exists(template_path) else Document()
    _clear_body(doc)
    vc, vbg = SIDE.get(R["verdict_side"], SIDE["none"])
    kc = KIND_COLOR.get(R["kind"], MUTED)

    # ---------------- page 1
    p = _para(doc, after=0)
    _run(p, R["title"], bold=True, size=15, color=NAVY)
    _run(p, "   " + R["kind_label"].upper(), bold=True, size=9, color=kc)
    p = _para(doc, after=6)
    bits = [f"Log file: {report.get('filename', '')}"]
    for k, lbl in (("msn", "MSN"), ("drone", "Model"), ("date", "Flight date"), ("reported_by", "Reported by")):
        if mission.get(k):
            bits.append(f"{lbl}: {mission[k]}")
    _run(p, " · ".join(bits), size=8, color=MUTED)

    _table(doc, [[{"runs": [{"t": "VERDICT:  ", "b": True, "sz": 11, "color": vc}, {"t": R["verdict_word"], "b": True, "sz": 15, "color": vc},
                            {"t": "   " + R["verdict_sub"], "sz": 8.5, "color": "3A4A54"}], "fill": vbg}]], [6.5])
    _section(doc, "In simple words")
    _run(_para(doc, after=4), R["simple"], size=10)
    _table(doc, [[{"t": k, "b": True, "fill": "EAF6FB"}, {"t": v, "b": i == 0}] for i, (k, v) in enumerate(R["summary"])], [1.25, 5.25])

    if R["wrong"]:
        _section(doc, R["wrong_title"])
        for i, w in enumerate(R["wrong"]):
            p = _para(doc, after=1.5, indent=0.28, hang=0.28)
            _run(p, f"{i + 1}.  ", bold=True)
            if w.get("t"):
                _run(p, w["t"] + "  ", bold=True, color=TEAL)
            _run(p, w["text"])
    if R["causes"]:
        _section(doc, "Possible causes")
        _table(doc, [[{"t": c["cat"], "b": True, "color": SIDE["human" if c["side"] == "Human" else "system"][0]},
                      {"t": c["role"], "b": c["role"] == "cause", "color": ROLE_COLOR[c["role"]]},
                      {"runs": [{"t": c["title"] + ". ", "b": True}, {"t": c["why"]}]}] for c in R["causes"]],
               [1.2, 1.0, 4.3], head=["Error type", "Role", "Cause and evidence"])
        _run(_para(doc, after=2), "Cause = made the outcome happen · Trigger = started the chain of events · Contributing = made it worse.", italic=True, size=7.5, color=GREY)
    _section(doc, "Why exactly it went wrong")
    _run(_para(doc, after=3), R["why"], size=10)
    if R["ruled"]:
        _run(_para(doc, after=3), R["ruled"], size=8.5, color="3A4A54")
    if R["actions"]:
        _section(doc, "Corrective actions")
        for a in R["actions"]:
            p = _para(doc, after=1, indent=0.25, hang=0.2)
            _run(p, "•  ", bold=True, color=TEAL)
            _run(p, a)

    # ---------------- annexure
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _run(_para(doc, after=1), "Annexure – complete flight data", bold=True, size=13, color=NAVY)
    _run(_para(doc, after=4), 'All values are read from the drone\'s flight log. "Healthy range" shows what a normal flight looks like; values marked CHECK are outside it.', size=8.5, color=MUTED)
    _section(doc, "A. Flight figures")
    rows = []
    for k, v, rng, mean, st in R["data"]:
        chk = st == "check"
        rows.append([{"t": k, "b": True}, {"t": v + ("  CHECK" if chk else ""), "b": chk, "color": "B3261E" if chk else None}, {"t": rng, "color": MUTED}, {"t": mean, "color": "3A4A54"}])
    _table(doc, rows, [1.35, 1.6, 1.15, 2.4], head=["Measurement", "Value", "Healthy range", "What it means"], size=8.5)

    with tempfile.TemporaryDirectory() as tmp:
        _section(doc, "B. Flight path")
        p3 = os.path.join(tmp, "path3d.png")
        if _path_3d(inv, p3):
            doc.add_picture(p3, width=Inches(6.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(_para(doc, after=4, align=WD_ALIGN_PARAGRAPH.CENTER), "The full flight path in 3D, coloured by height (dark low, bright high). Thin posts drop to the ground; the grey line is its shadow.", italic=True, size=7.5, color=MUTED)
        pt = os.path.join(tmp, "pathtop.png")
        if _path_top(inv, pt):
            doc.add_picture(pt, width=Inches(3.6))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(_para(doc, after=4, align=WD_ALIGN_PARAGRAPH.CENTER), "Seen from above, north up, coloured by height.", italic=True, size=7.5, color=MUTED)
        _table(doc, [[{"t": l["name"], "b": True}, f"{l['lat']:.7f}, {l['lon']:.7f}", f"{l['alt']} m" if l["alt"] is not None else "–", l["t"]] for l in inv.get("locations") or []],
               [1.6, 2.2, 1.2, 1.5], head=["Point", "Latitude, longitude", "Altitude (sea level)", "Flight time"], size=8.5)

        _section(doc, "C. Timeline of events")
        _table(doc, [[w["t"], w["clock"], {"t": w["text"], "b": w["sev"] == "crit", "color": "B3261E" if w["sev"] == "crit" else "9A5800" if w["sev"] == "warn" else None}] for w in R["timeline"]],
               [0.9, 0.9, 4.7], head=["Flight time", "Clock", "Event"], size=8.5)
        _run(_para(doc, after=2), "Flight time is counted from the moment the motors were armed (T+).", italic=True, size=7.5, color=GREY)

        _section(doc, "D. Health checks")
        order = {"crit": 0, "warn": 1, "info": 2, "ok": 3}
        _table(doc, [[{"t": SEV[c["sev"]][0], "b": True, "color": SEV[c["sev"]][1]}, {"t": c["title"], "b": True}, c["text"]] for c in sorted(inv.get("checks") or [], key=lambda c: order[c["sev"]])],
               [0.8, 1.9, 3.8], head=["Result", "Check", "Detail"], size=8.5)

        ch = inv.get("chart") or {}
        if ch.get("t"):
            _section(doc, "E. Flight data charts")
            specs = [
                ("height", [("Height", ch.get("height"), {"color": "#0f6f7d"}), ("Autopilot target", ch.get("target"), {"color": "#333", "linestyle": "--", "linewidth": 0.8})], "Height (m)", None,
                 "Height above take-off and the height the autopilot was aiming for (dashed). Red shading: after the crash. Dotted lines: failsafes."),
                ("speed", [("Ground speed", ch.get("speed"), {"color": "#0f6f7d"}), ("Climb rate", ch.get("climb"), {"color": "#d9822b"})], "m/s", None,
                 "Ground speed and climb rate. A negative climb rate means descending."),
                ("motors", [(f"M{i + 1}", m, {}) for i, m in enumerate(ch.get("motors") or [])], "Motor output (%)", None,
                 "Power sent to each motor. They should move together; one motor pinned near 100% while the others drop points to a motor or propeller failure."),
                ("volt", [("Voltage", ch.get("volt"), {"color": "#0f6f7d"})], "Battery (V)", ch.get("cell_floor"),
                 "Battery voltage. The dashed red line is 3.5 V per cell, the safe landing level."),
                ("roll", [("Roll", ch.get("roll"), {"color": "#0f6f7d"}), ("Target", ch.get("droll"), {"color": "#333", "linestyle": "--", "linewidth": 0.8})], "Roll (°)", None,
                 "Actual roll against the autopilot's target. If they separate before any contact, the drone was not doing what it was told: a sign of a system failure."),
            ]
            for name, ser, yl, floor, note in specs:
                png = os.path.join(tmp, f"c_{name}.png")
                if _chart(png, ch["t"], ser, yl, inv, floor):
                    doc.add_picture(png, width=Inches(6.3))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _run(_para(doc, after=4, align=WD_ALIGN_PARAGRAPH.CENTER), note, italic=True, size=7.5, color=MUTED)

    params = report.get("params") or {}
    srows = []
    for key, what, fmt in SETTINGS:
        if key in params:
            try:
                val = fmt(float(params[key]))
            except (TypeError, ValueError):
                val = str(params[key])
            srows.append([what, {"t": val, "b": True}, {"t": key, "color": MUTED}])
    if srows:
        _section(doc, "F. Safety settings on the drone")
        _table(doc, srows, [2.6, 2.3, 1.6], head=["What it controls", "Setting", "Parameter"], size=8.5)

    _run(_para(doc, after=0, before=6),
         f"Generated {datetime.now().strftime('%d %b %Y')} from the drone's flight log by the Neosky Drone Diagnostic Platform. "
         "How the verdict is reached: each finding is marked as human (pilot or setup) or system (aircraft) and weighted by whether it caused the "
         "outcome, started it or made it worse. A system error is ruled out only where the log shows that part working. This is an "
         "evidence-based assessment, not a certified accident investigation.", italic=True, size=7.5, color=GREY)
    doc.save(output_path)
    return output_path
